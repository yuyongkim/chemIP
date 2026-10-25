"""Spreadsheet helpers shared by the collectors that read official .xlsx downloads (CA ECCC, GB HSE ...).

openpyxl lives in the S: venv only. Workbooks are opened read-only with cached values (no formulas run).
"""
from __future__ import annotations

import re
from pathlib import Path

from intl_common import cas_tokens, clean


def sheet_rows(path: Path, sheet: str | int = 0) -> list[tuple]:
    """All rows of one sheet as tuples of cell values. sheet = name or position."""
    import openpyxl  # S: venv only

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def sheet_names(path: Path) -> list[str]:
    import openpyxl  # S: venv only

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def cell(value) -> str:
    """Cell value as clean text. Excel's '_x000D_' carriage-return escape is dropped."""
    return clean(str(value).replace("_x000D_", " ")) if value is not None else ""


def cas_in_cell(value) -> list[str]:
    """CAS numbers printed in a cell. Soft hyphens (U+00AD, invisible) and a space typed inside a number
    ('11138- 47-9') are closed up first; the digits themselves are never changed."""
    text = cell(value).replace("\u00ad ", "").replace("\u00ad", "")  # soft hyphen = invisible line-break hint
    text = re.sub(r"(\d)\s*-\s+(\d)", r"\1-\2", text)
    text = re.sub(r"(\d)\s+-\s*(\d)", r"\1-\2", text)
    return cas_tokens(text)
