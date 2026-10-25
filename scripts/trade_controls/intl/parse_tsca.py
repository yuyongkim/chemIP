"""US EPA TSCA lists.

- TSCA Inventory (non-confidential portion), EPA CSV download: CASRN, ChemName, ACTIVITY (ACTIVE/INACTIVE), FLAG.
  Flag meanings are copied from EPA's "How to Access the TSCA Inventory" page (sources/us_tsca.json).
- TSCA section 12(b) export notification list, EPA spreadsheet (sheets '12(b) Chemicals' and
  'SNURed Inactive PFAS chemicals'). Rows without a CAS (PMN / accession numbers only) are counted, not stored.
- TSCA section 6 risk-management substances: inventory rows flagged R + 12(b) rows whose Section is 6 / 6(a) / 6(h).
"""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

from intl_common import cas_tokens, clean, entry, fetch, file_info, split_entries

REGULATORY_FLAGS = ("5E", "5F", "12C", "R", "S", "SP", "T", "TP")


def _flags(raw: str) -> list[str]:
    return [f.strip() for f in (raw or "").split(";") if f.strip()]


def read_inventory(path: Path) -> tuple[list[dict], str]:
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.upper().startswith("TSCAINV_") and n.lower().endswith(".csv"))
        text = z.read(name).decode("utf-8", errors="replace")
    return list(csv.DictReader(io.StringIO(text))), name


def inventory(rows: list[dict], csv_name: str, flag_meaning: dict, files: list[dict]) -> dict:
    out = []
    for r in rows:
        flags = _flags(r.get("FLAG", ""))
        reg = [f"{f}({flag_meaning.get(f, '')})" for f in flags if f in REGULATORY_FLAGS]
        other = [f for f in flags if f not in REGULATORY_FLAGS]
        detail = "; ".join(p for p in (r.get("ACTIVITY", ""), ", ".join(reg), ", ".join(other),
                                        "UVCB" if r.get("UVCB") else "", f"EXP {r['EXP']}" if r.get("EXP") else "") if p)
        out.append(entry("", r.get("ChemName", "")[:300], detail, cas_all=[clean(r.get("CASRN", ""))]))
    entries, families, bad = split_entries(out)
    active = sum(1 for r in rows if r.get("ACTIVITY") == "ACTIVE")
    return {"entries": entries, "family_entries": families, "problems": [f"invalid CAS in source: {bad[:20]}"] if bad else [],
            "official_count": len(rows), "official_count_basis": f"{csv_name} 행 수(비CBI 목록 전체)",
            "counted": len(entries) + len(families),
            "source_version": f"Non-confidential TSCA Inventory {csv_name.split('_')[-1].split('.')[0]} (ACTIVE {active} / INACTIVE {len(rows) - active})",
            "notes": ["CBI(기밀) 물질 약 1.6만 건은 CAS 비공개라 이 목록에 없다(PMNACC 표, 일반명만 공개).",
                      "규제 플래그 뜻: " + "; ".join(f"{k}={v}" for k, v in flag_meaning.items())],
            "files": files}


def read_12b(path: Path) -> tuple[list[tuple], list[tuple]]:
    import openpyxl  # S: venv only

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    main = list(wb["12(b) Chemicals"].iter_rows(values_only=True))
    pfas = list(wb["SNURed Inactive PFAS chemicals"].iter_rows(values_only=True))
    return main, pfas


def _s(v) -> str:
    return clean("" if v is None else str(v))


def export_12b(main: list[tuple], pfas: list[tuple], files: list[dict], version: str) -> tuple[dict, list[dict]]:
    head = [_s(h) for h in main[0]]
    col = {h: i for i, h in enumerate(head)}
    rows, sec6 = [], []
    no_cas = 0
    for r in main[1:]:
        if not any(r):
            continue
        cas, section = _s(r[col["CAS_NUMBER"]]), _s(r[col["Section"]])
        detail = "; ".join(p for p in (f"TSCA §{section}" if section else "", _s(r[col["CFR_CITATION"]]),
                                        _s(r[col["FR_NOTICE_INFORMATION"]]), _s(r[col["CHEMICAL_CATEGORY"]]),
                                        f"END {_s(r[col['END DATE']])}" if r[col["END DATE"]] else "") if p)
        name = _s(r[col["CHEMICAL_NAME"]]) or _s(r[col["TSCA §12(b) Notification Name"]])
        if not cas:
            no_cas += 1
            continue
        item = entry("", name, detail, cas_all=cas_tokens(cas))
        rows.append(item)
        if section.startswith("6"):
            sec6.append(entry("", name, detail, cas_all=cas_tokens(cas), source="12(b) list"))
    for r in pfas[1:]:
        cas = _s(r[0])
        if cas:
            rows.append(entry("", _s(r[1]) or _s(r[4]), f"Inactive PFAS SNUR (40 CFR 721.11777); {_s(r[5])}", cas_all=cas_tokens(cas)))
        elif any(r):
            no_cas += 1
    counted = len(rows) + no_cas
    official = sum(1 for r in main[1:] if any(r)) + sum(1 for r in pfas[1:] if any(r))
    entries, families, bad = split_entries(rows)
    return ({"entries": entries, "family_entries": families,
             "problems": [f"invalid CAS in source: {bad}"] if bad else [],
             "official_count": official, "official_count_basis": "EPA 12(b) 스프레드시트 두 시트의 데이터 행 수",
             "counted": counted, "source_version": version,
             "notes": [f"CAS 없이 PMN·접수번호로만 적힌 {no_cas}행은 CAS 조회가 불가능해 저장하지 않음",
                       "EPA 고지: 이 목록은 참고용이며 누락이 있을 수 있다. 목록에 없다고 12(b) 면제는 아니다(Federal Register 확인 의무)."],
             "files": files}, sec6)


def section6(inv_rows: list[dict], sec6_12b: list[dict], files: list[dict], version: str) -> dict:
    rows = list(sec6_12b)
    for r in inv_rows:
        if "R" in _flags(r.get("FLAG", "")):
            rows.append(entry("", r.get("ChemName", ""), "TSCA Inventory flag R (section 6 risk management rule, proposed or final)",
                              cas_all=[clean(r.get("CASRN", ""))], source="inventory flag R"))
    entries, families, bad = split_entries(rows)
    return {"entries": entries, "family_entries": families, "problems": [f"invalid CAS in source: {bad}"] if bad else [],
            "official_count": None, "official_count_basis": "EPA 단일 공식 목록 없음 — Inventory R 플래그 + 12(b) 목록 §6 행의 합집합",
            "counted": len(rows), "source_version": version,
            "notes": ["두 EPA 자료의 갱신 시점이 달라 가장 최근 §6 규칙 일부는 빠질 수 있다(Federal Register 확인)."],
            "files": files}


def collect(spec: dict, raw_dir: Path, offline: bool) -> dict:
    d = spec["downloads"]
    inv_path = fetch(d["inventory"]["url"], raw_dir / d["inventory"]["file"], offline)
    b_path = fetch(d["export_12b"]["url"], raw_dir / d["export_12b"]["file"], offline)
    inv_info, b_info = file_info(inv_path, d["inventory"]["url"]), file_info(b_path, d["export_12b"]["url"])
    inv_rows, csv_name = read_inventory(inv_path)
    main, pfas = read_12b(b_path)
    b_version = f"EPA TSCA 12(b) chemical list {d['export_12b']['version']}"
    export, sec6_rows = export_12b(main, pfas, [b_info], b_version)
    return {
        "inventory": inventory(inv_rows, csv_name, spec["flag_meaning"], [inv_info]),
        "export_12b": export,
        "section6": section6(inv_rows, sec6_rows, [inv_info, b_info], f"{csv_name} flags + {b_version}"),
    }
