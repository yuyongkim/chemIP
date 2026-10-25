"""Load PRTR aggregate workbook sheets into the PRTR-only SQLite cache."""

from __future__ import annotations

import argparse
import re
import sqlite3
from pathlib import Path
from typing import Any

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
PRTR_DATA_DIR = PROJECT_ROOT / "data" / "prtr"
DEFAULT_DB_PATH = PRTR_DATA_DIR / "prtr_facilities.sqlite"
PROTECTED_DB_NAMES = {"terminology.db", "global_patent_index.db", "uspto_index.db"}


def compact(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_cas(value: str) -> str:
    digits = re.sub(r"\D", "", value or "")
    return digits.lstrip("0") if len(digits) >= 3 else ""


def to_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def validate_output_db_path(db_path: Path) -> Path:
    resolved = db_path.expanduser().resolve(strict=False)
    data_dir = PRTR_DATA_DIR.resolve(strict=False)
    if resolved.name.lower() in PROTECTED_DB_NAMES:
        raise SystemExit(f"Refusing to write PRTR aggregates to protected DB: {resolved}")
    if not resolved.is_relative_to(data_dir):
        raise SystemExit(f"Refusing to write PRTR aggregates outside {data_dir}: {resolved}")
    if "prtr" not in resolved.stem.lower():
        raise SystemExit(f"Refusing ambiguous PRTR cache filename: {resolved}")
    return resolved


def find_workbook() -> Path:
    matches = sorted(PROJECT_ROOT.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    for path in matches:
        if "조사결과" in path.name:
            return path
    raise SystemExit("Could not find 조사결과 workbook. Pass --workbook explicitly.")


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS prtr_aggregate_materials (
            year INTEGER NOT NULL,
            cas_no TEXT NOT NULL,
            cas_key TEXT NOT NULL,
            chemical_name TEXT NOT NULL,
            release_total_kg REAL NOT NULL DEFAULT 0,
            air_release_kg REAL NOT NULL DEFAULT 0,
            water_release_kg REAL NOT NULL DEFAULT 0,
            soil_release_kg REAL NOT NULL DEFAULT 0,
            self_landfill_kg REAL NOT NULL DEFAULT 0,
            transfer_total_kg REAL NOT NULL DEFAULT 0,
            wastewater_transfer_kg REAL NOT NULL DEFAULT 0,
            waste_transfer_kg REAL NOT NULL DEFAULT 0,
            reported_total_kg REAL NOT NULL DEFAULT 0,
            source_file TEXT NOT NULL DEFAULT '',
            PRIMARY KEY (year, cas_key, chemical_name)
        );

        CREATE TABLE IF NOT EXISTS prtr_aggregate_breakdowns (
            year INTEGER NOT NULL,
            dimension_type TEXT NOT NULL,
            dimension_name TEXT NOT NULL,
            cas_no TEXT NOT NULL,
            cas_key TEXT NOT NULL,
            chemical_name TEXT NOT NULL,
            release_total_kg REAL NOT NULL DEFAULT 0,
            air_release_kg REAL NOT NULL DEFAULT 0,
            water_release_kg REAL NOT NULL DEFAULT 0,
            soil_release_kg REAL NOT NULL DEFAULT 0,
            self_landfill_kg REAL NOT NULL DEFAULT 0,
            transfer_total_kg REAL NOT NULL DEFAULT 0,
            wastewater_transfer_kg REAL NOT NULL DEFAULT 0,
            waste_transfer_kg REAL NOT NULL DEFAULT 0,
            reported_total_kg REAL NOT NULL DEFAULT 0,
            source_file TEXT NOT NULL DEFAULT '',
            PRIMARY KEY (year, dimension_type, dimension_name, cas_key, chemical_name)
        );

        CREATE TABLE IF NOT EXISTS prtr_carcinogenic_classes (
            year INTEGER NOT NULL,
            cas_no TEXT NOT NULL,
            cas_key TEXT NOT NULL,
            chemical_name TEXT NOT NULL,
            iarc_group TEXT NOT NULL,
            release_total_kg REAL NOT NULL DEFAULT 0,
            air_release_kg REAL NOT NULL DEFAULT 0,
            water_release_kg REAL NOT NULL DEFAULT 0,
            soil_release_kg REAL NOT NULL DEFAULT 0,
            self_landfill_kg REAL NOT NULL DEFAULT 0,
            transfer_total_kg REAL NOT NULL DEFAULT 0,
            wastewater_transfer_kg REAL NOT NULL DEFAULT 0,
            waste_transfer_kg REAL NOT NULL DEFAULT 0,
            reported_total_kg REAL NOT NULL DEFAULT 0,
            source_file TEXT NOT NULL DEFAULT '',
            PRIMARY KEY (year, cas_key, chemical_name, iarc_group)
        );

        CREATE INDEX IF NOT EXISTS idx_prtr_agg_material_cas ON prtr_aggregate_materials(year, cas_key);
        CREATE INDEX IF NOT EXISTS idx_prtr_agg_breakdown_cas ON prtr_aggregate_breakdowns(year, cas_key, dimension_type);
        CREATE INDEX IF NOT EXISTS idx_prtr_agg_carc_cas ON prtr_carcinogenic_classes(year, cas_key);
        """
    )
    conn.commit()


def metric_values(row: tuple[Any, ...], offset: int) -> tuple[float, ...]:
    release_total = to_number(row[offset])
    air = to_number(row[offset + 1])
    water = to_number(row[offset + 2])
    soil = to_number(row[offset + 3])
    landfill = to_number(row[offset + 4])
    transfer_total = to_number(row[offset + 5])
    wastewater = to_number(row[offset + 6])
    waste = to_number(row[offset + 7])
    reported_total = release_total + landfill + transfer_total
    return release_total, air, water, soil, landfill, transfer_total, wastewater, waste, reported_total


def load_material_sheet(conn: sqlite3.Connection, ws, *, year: int, source_file: str) -> int:
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        chemical_name = compact(row[1] if len(row) > 1 else "")
        cas_no = compact(row[2] if len(row) > 2 else "")
        if not chemical_name or chemical_name == "총합계":
            continue
        values = metric_values(row, 3)
        conn.execute(
            """
            INSERT OR REPLACE INTO prtr_aggregate_materials
            (year, cas_no, cas_key, chemical_name, release_total_kg, air_release_kg, water_release_kg,
             soil_release_kg, self_landfill_kg, transfer_total_kg, wastewater_transfer_kg,
             waste_transfer_kg, reported_total_kg, source_file)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (year, cas_no, normalize_cas(cas_no), chemical_name, *values, source_file),
        )
        count += 1
    return count


def load_breakdown_sheet(conn: sqlite3.Connection, ws, *, year: int, source_file: str, dimension_type: str) -> int:
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        dimension_name = compact(row[0] if len(row) > 0 else "")
        chemical_name = compact(row[1] if len(row) > 1 else "")
        cas_no = compact(row[2] if len(row) > 2 else "")
        if not dimension_name or not chemical_name or chemical_name == "총합계":
            continue
        values = metric_values(row, 3)
        conn.execute(
            """
            INSERT OR REPLACE INTO prtr_aggregate_breakdowns
            (year, dimension_type, dimension_name, cas_no, cas_key, chemical_name,
             release_total_kg, air_release_kg, water_release_kg, soil_release_kg,
             self_landfill_kg, transfer_total_kg, wastewater_transfer_kg, waste_transfer_kg,
             reported_total_kg, source_file)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (year, dimension_type, dimension_name, cas_no, normalize_cas(cas_no), chemical_name, *values, source_file),
        )
        count += 1
    return count


def load_carcinogenic_sheet(conn: sqlite3.Connection, ws, *, year: int, source_file: str) -> int:
    count = 0
    for row in ws.iter_rows(min_row=6, values_only=True):
        iarc_group = compact(row[1] if len(row) > 1 else "")
        cas_no = compact(row[2] if len(row) > 2 else "")
        chemical_name = compact(row[3] if len(row) > 3 else "")
        if not iarc_group or not chemical_name:
            continue
        values = metric_values(row, 4)
        conn.execute(
            """
            INSERT OR REPLACE INTO prtr_carcinogenic_classes
            (year, cas_no, cas_key, chemical_name, iarc_group, release_total_kg, air_release_kg,
             water_release_kg, soil_release_kg, self_landfill_kg, transfer_total_kg,
             wastewater_transfer_kg, waste_transfer_kg, reported_total_kg, source_file)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (year, cas_no, normalize_cas(cas_no), chemical_name, iarc_group, *values, source_file),
        )
        count += 1
    return count


def load(args: argparse.Namespace) -> None:
    db_path = validate_output_db_path(args.db)
    workbook_path = args.workbook or find_workbook()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    conn = sqlite3.connect(db_path)
    ensure_schema(conn)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    source_file = str(workbook_path)

    counts = {
        "materials": load_material_sheet(conn, wb["1. 화학물질별 배출량·위탁처리량"], year=args.year, source_file=source_file),
        "regions": load_breakdown_sheet(conn, wb["2. 지역별 화학물질 배출량·위탁처리량"], year=args.year, source_file=source_file, dimension_type="region"),
        "industrial_complexes": load_breakdown_sheet(conn, wb["3. 산단별 화학물질 배출량·위탁처리량"], year=args.year, source_file=source_file, dimension_type="industrial_complex"),
        "industries": load_breakdown_sheet(conn, wb["4. 업종별 화학물질 배출량·위탁처리량"], year=args.year, source_file=source_file, dimension_type="industry"),
        "carcinogenic_classes": load_carcinogenic_sheet(conn, wb["7. 발암우려물질 등급별 배출량·위탁처리량"], year=args.year, source_file=source_file),
    }
    conn.commit()
    for name, count in counts.items():
        print(f"{name}: {count}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, default=2024)
    parser.add_argument("--workbook", type=Path, default=None)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    return parser.parse_args()


if __name__ == "__main__":
    load(parse_args())
